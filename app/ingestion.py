"""Document ingestion: extract text from PDF, DOCX, XLSX, MD, TXT."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # pymupdf
from docx import Document as DocxDocument
from openpyxl import load_workbook


@dataclass
class TextBlock:
    text: str
    source: str  # e.g. "page 3", "sheet Sales", "paragraph 5"


@dataclass
class ExtractedDocument:
    filename: str
    format: str
    blocks: list[TextBlock] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks if b.text.strip())


def extract_pdf(data: bytes, filename: str) -> ExtractedDocument:
    doc = ExtractedDocument(filename=filename, format="pdf")
    try:
        pdf = fitz.open(stream=data, filetype="pdf")
        for i, page in enumerate(pdf):
            text = page.get_text()
            if text.strip():
                doc.blocks.append(TextBlock(text=text, source=f"page {i + 1}"))
            else:
                doc.warnings.append(f"Page {i + 1}: no extractable text (image-only?)")
        pdf.close()
    except Exception as e:
        doc.warnings.append(f"PDF extraction error: {e}")
    return doc


def extract_docx(data: bytes, filename: str) -> ExtractedDocument:
    doc = ExtractedDocument(filename=filename, format="docx")
    try:
        docx = DocxDocument(io.BytesIO(data))
        for i, para in enumerate(docx.paragraphs):
            if para.text.strip():
                doc.blocks.append(TextBlock(text=para.text, source=f"paragraph {i + 1}"))
        for table_idx, table in enumerate(docx.tables):
            rows_text = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows_text.append(" | ".join(cells))
            if any(r.strip() for r in rows_text):
                doc.blocks.append(TextBlock(
                    text="\n".join(rows_text),
                    source=f"table {table_idx + 1}",
                ))
    except Exception as e:
        doc.warnings.append(f"DOCX extraction error: {e}")
    return doc


def extract_xlsx(data: bytes, filename: str) -> ExtractedDocument:
    doc = ExtractedDocument(filename=filename, format="xlsx")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_text = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    rows_text.append(" | ".join(cells))
            if rows_text:
                doc.blocks.append(TextBlock(
                    text="\n".join(rows_text),
                    source=f"sheet '{sheet_name}'",
                ))
        wb.close()
    except Exception as e:
        doc.warnings.append(f"XLSX extraction error: {e}")
    return doc


def extract_text(data: bytes, filename: str) -> ExtractedDocument:
    doc = ExtractedDocument(filename=filename, format="txt")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
        doc.warnings.append("File decoded as latin-1 (not UTF-8)")
    doc.blocks.append(TextBlock(text=text, source="full document"))
    return doc


def extract_markdown(data: bytes, filename: str) -> ExtractedDocument:
    doc = ExtractedDocument(filename=filename, format="md")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
        doc.warnings.append("File decoded as latin-1 (not UTF-8)")
    doc.blocks.append(TextBlock(text=text, source="full document"))
    return doc


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".md": extract_markdown,
    ".txt": extract_text,
}


def ingest_document(data: bytes, filename: str) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    extractor = EXTRACTORS.get(suffix)
    if extractor is None:
        doc = ExtractedDocument(filename=filename, format=suffix)
        doc.warnings.append(f"Unsupported format: {suffix}")
        return doc
    return extractor(data, filename)
